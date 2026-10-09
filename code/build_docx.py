#!/usr/bin/env python
"""Build the three Word deliverables from filled markdown / letter text.

    python build_docx.py

Unresolved [[KEY]] placeholders stay visible so they cannot be mistaken
for typed numbers. Re-run after Stage 9.
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
MS_DIR = ROOT / "manuscript"
LETTER_DIR = ROOT / "letter"
DELIV = ROOT / "deliverables"

MARK_WORDS = (
    "withdrawn", "audit", "native unit", "groupby", "z-scored",
    "not reproduce", "Jaccard", "heatwave", "1 000 m", "Diebold",
    "observed-target", "[[",
)


def _ensure_docx():
    try:
        import docx  # noqa: F401
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx>=1.1"])


def _fill_sources():
    code = Path(__file__).resolve().parent
    for script in ("build_letter.py", "build_manuscript.py"):
        subprocess.check_call([sys.executable, str(code / script)])


def _add_para(doc, text: str, *, highlight: bool = False, bold: bool = False):
    from docx.enum.text import WD_COLOR_INDEX
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


_REV_INS = re.compile(r"==([^=]+)==")
_REV_DEL = re.compile(r"~~([^~]+)~~")


def _strip_revision_marks(text: str) -> str:
    """Journal clean copy: no markdown revision tokens."""
    text = _REV_DEL.sub("", text)
    text = _REV_INS.sub(r"\1", text)
    return text.replace("==", "").replace("~~", "")


def _add_rich_para(doc, text: str, *, marked: bool = False):
    """Marked copy: ==insert== underlined; ~~delete~~ struck through."""
    from docx.enum.text import WD_COLOR_INDEX, WD_UNDERLINE

    if not marked:
        return _add_para(doc, _strip_revision_marks(text))
    if "==" not in text and "~~" not in text:
        return _add_para(doc, text, highlight=_is_marked(text))
    p = doc.add_paragraph()
    pos = 0
    for m in re.finditer(r"==([^=]+)==|~~([^~]+)~~", text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        if m.group(1):
            r = p.add_run(m.group(1))
            r.underline = WD_UNDERLINE.SINGLE
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif m.group(2):
            r = p.add_run(m.group(2))
            r.font.strike = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])
    return p


def _body_line(doc, line: str, *, marked: bool, bullet: bool = False):
    if bullet:
        if marked and ("==" in line or "~~" in line):
            _add_rich_para(doc, line, marked=True)
        else:
            plain = _strip_revision_marks(line) if not marked else line
            p = doc.add_paragraph(plain, style="List Bullet")
            if marked and _is_marked(line):
                for run in p.runs:
                    from docx.enum.text import WD_COLOR_INDEX
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return
    if marked and ("==" in line or "~~" in line):
        _add_rich_para(doc, line, marked=True)
    else:
        plain = _strip_revision_marks(line) if not marked else line
        _add_para(doc, plain, highlight=marked and _is_marked(plain),
                  bold=line.startswith("**") and line.endswith("**"))


def _is_marked(text: str) -> bool:
    low = text.lower()
    return any(w.lower() in low for w in MARK_WORDS)


def markdown_to_docx(src: Path, dest: Path, marked: bool) -> None:
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    if marked:
        _add_para(doc,
                  "MARKED-UP REVISION — underline + yellow = new or rewritten text; "
                  "strikethrough = withdrawn claim. Compare against the submitted "
                  "manuscript PDF for full diff.",
                  highlight=True, bold=True)

    for raw in src.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(_strip_revision_marks(line[2:].strip()), level=0)
            continue
        if line.startswith("## "):
            doc.add_heading(_strip_revision_marks(line[3:].strip()), level=1)
            continue
        if line.startswith("### "):
            doc.add_heading(_strip_revision_marks(line[4:].strip()), level=2)
            continue
        if line.startswith("#### "):
            doc.add_heading(_strip_revision_marks(line[5:].strip()), level=3)
            continue
        if line.startswith("---"):
            continue
        if line.startswith("|") and line.endswith("|"):
            continue
        if line.startswith("- "):
            _body_line(doc, line[2:], marked=marked, bullet=True)
            continue
        _body_line(doc, line, marked=marked)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    print(f"wrote {dest}")


def letter_to_docx(src: Path, dest: Path, marked: bool) -> None:
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    if marked:
        _add_para(doc,
                  "Unresolved [[KEY]] values are highlighted. They are filled "
                  "only from outputs/manuscript_numbers.json after Stage 9.",
                  highlight=True, bold=True)
    for raw in src.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line:
            continue
        _add_para(doc, line, highlight=marked and _is_marked(line),
                  bold=line in ("Response to Reviewers", "Letter to the Editor",
                                "Dear Editor,", "Summary of changes"))
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    print(f"wrote {dest}")


def main():
    _ensure_docx()
    _fill_sources()
    filled_ms = MS_DIR / "manuscript_filled.md"
    if not filled_ms.exists():
        filled_ms = MS_DIR / "manuscript_master.md"
    filled_lt = LETTER_DIR / "response_letter_filled.txt"
    if not filled_lt.exists():
        filled_lt = LETTER_DIR / "letter_text.txt"

    clean = MS_DIR / "manuscript_clean.docx"
    marked = MS_DIR / "manuscript_marked_up.docx"
    letter = LETTER_DIR / "response_letter.docx"
    markdown_to_docx(filled_ms, clean, marked=False)
    markdown_to_docx(filled_ms, marked, marked=True)
    letter_to_docx(filled_lt, letter, marked=True)

    DELIV.mkdir(exist_ok=True)
    for src, name in ((clean, "manuscript_clean.docx"),
                      (marked, "manuscript_marked_up.docx"),
                      (letter, "response_letter.docx")):
        dest = DELIV / name
        dest.write_bytes(src.read_bytes())
        print(f"copied {dest}")

    leftover = re.findall(r"\[\[([A-Z0-9_@]+)\]\]", filled_ms.read_text(encoding="utf-8"))
    print(f"manuscript open placeholders: {len(set(leftover))}")
    leftover_l = re.findall(r"\[\[([A-Z0-9_@]+)\]\]", filled_lt.read_text(encoding="utf-8"))
    print(f"letter open placeholders: {len(set(leftover_l))}")
    try:
        subprocess.check_call([sys.executable, str(Path(__file__).parent / "verify_build.py")])
    except subprocess.CalledProcessError:
        print("warning: verify_build failed — fix before sending to client")


if __name__ == "__main__":
    main()
