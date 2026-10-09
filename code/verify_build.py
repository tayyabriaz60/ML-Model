"""Pre-send checks on built deliverables."""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DELIV = ROOT / "deliverables"
FIGS = Path(__file__).resolve().parent / "outputs" / "figures"
EXPECTED_FIG = {f"F{i}_" for i in range(1, 12)}


def _docx_text(path: Path) -> str:
    from docx import Document
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


def _marked_runs(path: Path) -> tuple[int, int, int]:
    from docx import Document
    from docx.enum.text import WD_UNDERLINE
    doc = Document(str(path))
    under = strike = 0
    literal_eq = 0
    for p in doc.paragraphs:
        if "==" in p.text or "~~" in p.text:
            literal_eq += 1
        for run in p.runs:
            if run.font.strike:
                strike += 1
            if run.font.underline and run.font.underline != WD_UNDERLINE.NONE:
                under += 1
    return under, strike, literal_eq


def main():
    errs = []
    clean = DELIV / "manuscript_clean.docx"
    marked = DELIV / "manuscript_marked_up.docx"
    if clean.exists():
        t = _docx_text(clean)
        if "==" in t or "~~" in t:
            errs.append(f"{clean.name}: literal == or ~~ in body ({t.count('==')} / {t.count('~~')})")
    if marked.exists():
        u, s, lit = _marked_runs(marked)
        if lit:
            errs.append(f"{marked.name}: {lit} paragraph(s) with literal == or ~~")
        if u == 0 and s == 0:
            print(f"warning: {marked.name}: no underlined or struck runs")
    if FIGS.is_dir():
        pngs = list(FIGS.glob("*.png"))
        names = [p.name for p in pngs]
        canonical = [n for n in names if any(n.startswith(pfx) for pfx in EXPECTED_FIG)]
        legacy = [n for n in names if n not in canonical]
        if legacy:
            errs.append(f"figures: legacy PNGs still present: {legacy[:5]}")
        if len(set(canonical)) != len([n for n in names if n in canonical]):
            pass
        # count unique figure ids F1..F11
        ids = set()
        for n in canonical:
            ids.add(n.split("_")[0])
        if len(ids) > 11:
            errs.append(f"figures: more than 11 figure ids in PNG names: {sorted(ids)}")
    if errs:
        print("VERIFY FAILED:")
        for e in errs:
            print(" ", e)
        sys.exit(1)
    print("verify_build: OK")


if __name__ == "__main__":
    main()
